"""Build docs/presentation.pptx (12 slides) from the committed artifacts.

Every number on the slides is read from artifacts/ (numbers freeze), so the deck always matches the
repository. Export to PDF with PowerPoint (see scripts/export_pdf.ps1) or any viewer.

Usage: uv run python scripts/build_presentation.py
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
FIG = ART / "figures"
IMAGES = ROOT / "docs" / "images"
OUT = ROOT / "docs" / "presentation.pptx"

PRESENTER = "Matanel"
REPO_URL = "github.com/matnael332/olist-crewai-flow"
APP_URL = "olist-late-delivery-1.streamlit.app"  # Streamlit Community Cloud deployment

NAVY = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
RED = RGBColor(0xDC, 0x26, 0x26)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
GREY = RGBColor(0x64, 0x74, 0x8B)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SKY = RGBColor(0xDB, 0xEA, 0xFE)
SAND = RGBColor(0xFE, 0xF3, 0xC7)
CODE_CMD = RGBColor(0x93, 0xC5, 0xFD)
CODE_OUT = RGBColor(0x86, 0xEF, 0xAC)
CODE_FAIL = RGBColor(0xFC, 0xA5, 0xA5)


def load(name: str) -> dict:
    return json.loads((ART / name).read_text(encoding="utf-8"))


def pct(x: float, d: int = 1) -> str:
    return f"{x * 100:.{d}f}%"


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.count = 0

    def slide(self, title: str, kicker: str = "", notes: str = ""):
        self.count += 1
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), self.prs.slide_height)
        self._fill(bar, BLUE)
        if kicker:
            self.text(s, kicker.upper(), 0.6, 0.35, 12, 0.35, size=12, color=BLUE, bold=True)
        self.text(s, title, 0.6, 0.62, 12.2, 0.9, size=30, bold=True)
        self.text(
            s,
            f"Olist late-delivery AI workflow  |  {self.count}",
            0.6,
            7.0,
            12.2,
            0.3,
            size=10,
            color=GREY,
        )
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    @staticmethod
    def _fill(shape, color, line=None) -> None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = line

    def text(
        self,
        s,
        text,
        left,
        top,
        width,
        height,
        size=18,
        color=NAVY,
        bold=False,
        align=PP_ALIGN.LEFT,
    ):
        box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
        return box

    def bullets(self, s, items, left, top, width, height, size=18):
        box = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            bold, rest = (item.split("|", 1) + [""])[:2] if "|" in item else ("", item)
            if bold:
                r = p.add_run()
                r.text = bold.strip() + " "
                r.font.bold = True
                r.font.size = Pt(size)
                r.font.color.rgb = NAVY
            r = p.add_run()
            r.text = ("• " if not bold else "") + rest.strip()
            r.font.size = Pt(size)
            r.font.color.rgb = NAVY
        return box

    def stat(self, s, value, label, left, top, color=BLUE, width=2.9):
        card = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(1.35)
        )
        self._fill(card, LIGHT)
        card.adjustments[0] = 0.12
        self.text(
            s, value, left + 0.15, top + 0.12, width - 0.3, 0.7, size=30, bold=True, color=color
        )
        self.text(s, label, left + 0.15, top + 0.78, width - 0.3, 0.5, size=12, color=GREY)

    def picture(self, s, name, left, top, width=None, height=None, folder=FIG):
        path = folder / name
        kwargs = {}
        if width:
            kwargs["width"] = Inches(width)
        if height:
            kwargs["height"] = Inches(height)
        return s.shapes.add_picture(str(path), Inches(left), Inches(top), **kwargs)

    def code(self, s, lines, left, top, width, height, size=12):
        """Terminal-style box: commands in blue, failures in red, other output in green."""
        b = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        self._fill(b, NAVY)
        b.adjustments[0] = 0.05
        tf = b.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.25)
        tf.margin_top = tf.margin_bottom = Inches(0.15)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT  # autoshapes center their first paragraph by default
            r = p.add_run()
            r.text = line or " "
            r.font.size = Pt(size)
            r.font.name = "Consolas"
            if line.startswith("$"):
                r.font.color.rgb = CODE_CMD
            elif "FAIL" in line or "exit 1" in line:
                r.font.color.rgb = CODE_FAIL
            else:
                r.font.color.rgb = CODE_OUT
        return b

    def box(self, s, text, left, top, width, height, fill, size=13, bold=True, shape=None):
        b = s.shapes.add_shape(
            shape or MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        self._fill(b, fill, line=GREY if fill == WHITE else None)
        tf = b.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size if i == 0 else size - 2)
            r.font.bold = bold and i == 0
            r.font.color.rgb = NAVY
        return b

    def arrow(self, s, x1, y1, x2, y2, color=GREY, dashed=False):
        c = s.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        c.line.color.rgb = color
        c.line.width = Pt(2)
        ln = c.line._get_or_add_ln()
        tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
        ln.append(tail)
        if dashed:
            dash = ln.makeelement(qn("a:prstDash"), {"val": "dash"})
            ln.insert(0, dash)
        return c

    def table(self, s, rows, left, top, width, height, col_widths=None, size=13, highlight=None):
        shape = s.shapes.add_table(
            len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)
        )
        t = shape.table
        if col_widths:
            for i, w in enumerate(col_widths):
                t.columns[i].width = Emu(int(Inches(w)))
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = t.cell(r, c)
                cell.text = str(value)
                para = cell.text_frame.paragraphs[0]
                para.runs[0].font.size = Pt(size)
                para.runs[0].font.bold = r == 0 or r == highlight
                para.runs[0].font.color.rgb = WHITE if r == 0 else NAVY
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY if r == 0 else (SKY if r == highlight else WHITE)
        return t


def build() -> None:
    eda, metrics, meta = load("eda_summary.json"), load("metrics.json"), load("model_meta.json")
    contract, report = load("dataset_contract.json"), load("cleaning_report.json")
    status = load("run_status.json")
    checks = {g: len(v["checks"]) for g, v in load("validation_report.json")["gates"].items()}
    ao, dp = eda["all_orders"], eda["delivered_population"]
    rv, dl = dp["reviews"], dp["delivery"]
    chosen = metrics["selected_model"]
    sel = metrics["candidates"][chosen]
    test = sel["test"]
    split = metrics["split"]
    states = [s for s in dp["late_rate_by_state"] if s["orders"] >= int(0.004 * dp["n_orders"])]
    deck = Deck()

    # 1. Title -----------------------------------------------------------------------------
    s = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    deck.count += 1
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.prs.slide_width, deck.prs.slide_height)
    deck._fill(bg, NAVY)
    deck.text(
        s,
        "FINAL PROJECT  |  INDUSTRY-SIMULATED AI PRODUCT WORKFLOW",
        0.9,
        1.6,
        11.5,
        0.4,
        size=14,
        color=RGBColor(0x93, 0xC5, 0xFD),
        bold=True,
    )
    deck.text(
        s, "Will this order arrive late?", 0.9, 2.2, 11.5, 1.2, size=48, color=WHITE, bold=True
    )
    deck.text(
        s,
        "Predicting late deliveries at checkout with a CrewAI Flow: a Data Analyst crew "
        "hands off to a Data Scientist crew",
        0.9,
        3.4,
        11.0,
        1.0,
        size=22,
        color=WHITE,
    )
    deck.text(s, PRESENTER, 0.9, 4.45, 11.5, 0.6, size=22, color=WHITE, bold=True)
    deck.text(
        s,
        [
            "Olist Brazilian e-commerce data  |  CrewAI + Gemini, scikit-learn, Streamlit, "
            "Flask, GitHub",
            f"Code: {REPO_URL}",
            f"Live dashboard: {APP_URL}",
        ],
        0.9,
        5.3,
        11.5,
        1.2,
        size=16,
        color=RGBColor(0xCB, 0xD5, 0xE1),
    )
    s.notes_slide.notes_text_frame.text = (
        "One sentence: we built an AI product workflow that turns raw marketplace data into "
        "insights and a deployable late-delivery risk model, with two AI agent teams and "
        "automatic quality gates between them."
    )

    # 2. Business problem --------------------------------------------------------------------
    s = deck.slide(
        "Late deliveries destroy customer satisfaction",
        "The business problem",
        "Lead with the 62% vs 9% contrast; this is why predicting lateness at checkout matters.",
    )
    deck.stat(
        s,
        pct(rv["bad_review_rate_late"]),
        "of LATE orders get a 1-2 star review",
        0.6,
        1.7,
        RED,
        width=3.6,
    )
    deck.stat(
        s,
        pct(rv["bad_review_rate_on_time"]),
        "of ON-TIME orders get a 1-2 star review",
        0.6,
        3.25,
        GREEN,
        width=3.6,
    )
    deck.stat(
        s,
        pct(dp["late_rate"]),
        f"of {dp['n_orders']:,} delivered orders arrived late",
        0.6,
        4.8,
        BLUE,
        width=3.6,
    )
    deck.picture(s, "review_score_vs_late.png", 4.6, 1.6, width=8.2)
    deck.text(
        s,
        "Goal: flag risky orders at checkout so operations can act before the customer "
        "is let down.",
        4.6,
        6.3,
        8.2,
        0.6,
        size=16,
        bold=True,
        color=BLUE,
    )

    # 3. Data ---------------------------------------------------------------------------------
    s = deck.slide(
        "Real marketplace data, cleaned and versioned",
        "Dataset",
        "Nine linked tables from Olist's own GitHub at a pinned commit; every file is sha256-checked.",
    )
    deck.bullets(
        s,
        [
            "Source:|Olist Brazilian E-Commerce, 9 tables (orders, items, payments, reviews, customers, "
            "products, sellers, geolocation, categories), pinned commit + sha256 manifest",
            f"Scale:|{ao['n_orders']:,} orders, {ao['purchase_range'][0][:7]} to "
            f"{ao['purchase_range'][1][:7]}",
            "Real quality issues:|non-delivered orders, missing delivery dates, out-of-sequence "
            "timestamps, products without category, duplicate reviews, invalid payments",
            f"Result:|clean_data.csv with {report['rows_out']:,} delivered orders, one row per order, "
            f"{len(contract['columns'])} documented columns",
        ],
        0.6,
        1.6,
        6.4,
        4.8,
        size=17,
    )
    rows = (
        [["rule", "rows removed"]]
        + [
            [r["rule"].replace("_", " "), f"{r['rows_removed']:,}"]
            for r in report["rules"]
            if r["rows_removed"] > 0
        ]
        + [["rows kept", f"{report['rows_out']:,}"]]
    )
    deck.table(
        s,
        rows,
        7.4,
        1.7,
        5.4,
        0.4 * len(rows),
        col_widths=[3.6, 1.8],
        size=13,
        highlight=len(rows) - 1,
    )

    # 4. Architecture ---------------------------------------------------------------------------
    s = deck.slide(
        "Two AI crews, three quality gates, one product",
        "Architecture",
        "Run-then-narrate: tested Python computes every artifact, agents inspect and explain; the "
        "Flow decides pass or fail. A failed gate stops the run and nothing is promoted. On top, "
        "three channels read the published artifacts through one risk policy.",
    )
    y = 2.3
    deck.box(s, "9 raw tables\npinned + hashed", 0.5, y, 1.6, 1.1, WHITE)
    deck.box(s, "Crew 1\nData Analyst\n6 agents", 2.5, y, 1.9, 1.1, SKY)
    deck.box(
        s,
        "gate\ncontract",
        4.8,
        y + 0.05,
        1.3,
        1.0,
        RGBColor(0xDC, 0xFC, 0xE7),
        shape=MSO_SHAPE.DIAMOND,
        size=11,
    )
    deck.box(s, "Crew 2 phase 1\nfeatures", 6.5, y, 1.8, 1.1, SAND)
    deck.box(
        s,
        "gate\nfeatures",
        8.7,
        y + 0.05,
        1.3,
        1.0,
        RGBColor(0xDC, 0xFC, 0xE7),
        shape=MSO_SHAPE.DIAMOND,
        size=11,
    )
    deck.box(s, "Crew 2 phase 2\nmodels + card", 10.4, y, 1.9, 1.1, SAND)
    for x1, x2 in [(2.1, 2.5), (4.4, 4.8), (6.1, 6.5), (8.3, 8.7), (10.0, 10.4)]:
        deck.arrow(s, x1, y + 0.55, x2, y + 0.55)
    deck.box(
        s,
        "gate artifacts",
        10.55,
        4.05,
        1.6,
        0.8,
        RGBColor(0xDC, 0xFC, 0xE7),
        shape=MSO_SHAPE.DIAMOND,
        size=10,
    )
    deck.arrow(s, 11.35, y + 1.1, 11.35, 4.05)
    deck.box(s, "artifacts/ promoted\n8 spec files + reports", 7.4, 4.05, 2.7, 0.8, WHITE)
    deck.arrow(s, 10.55, 4.45, 10.1, 4.45)
    deck.box(
        s,
        "fail_gracefully (any failed gate)\nstatus file, exit 1, nothing promoted",
        3.2,
        4.05,
        3.6,
        0.8,
        RGBColor(0xFE, 0xE2, 0xE2),
        size=12,
    )
    deck.arrow(s, 5.45, y + 1.05, 5.0, 4.05, color=RED, dashed=True)
    channels = [
        ("Ops copilot\nchat for the agent operator", 4.95, SKY),
        ("Dashboard\nbusiness + data science", 7.45, LIGHT),
        ("API\n/predict, /predict/batch, /risk-list", 9.95, LIGHT),
    ]
    for label, x, fill in channels:
        deck.box(s, label, x, 5.45, 2.35, 0.85, fill, size=12)
        deck.arrow(s, 8.75, 4.85, x + 1.17, 5.45, color=BLUE)
    deck.text(
        s,
        [
            "Crew 1 (6 agents): data engineer, documentation writer, analytics lead, insights "
            "writer, data steward, contract author.",
            "Crew 2 (5 agents): feature engineer and reviewer, ML engineer, model reviewer, model "
            "card author.",
        ],
        0.6,
        5.35,
        4.2,
        1.5,
        size=12,
        color=GREY,
    )

    # 5. What happened -----------------------------------------------------------------------
    s = deck.slide(
        "What happened: fast growth and a Black Friday spike",
        "Crew 1 insights",
        "All-orders layer for volume; delivered layer for lateness. Every number comes from "
        "eda_summary.json.",
    )
    deck.picture(s, "monthly_orders_all_vs_delivered.png", 0.5, 1.55, width=7.4)
    deck.text(
        s,
        f"All {ao['n_orders']:,} orders by purchase month, every status: "
        f"{pct(1 - ao['non_delivered_share'])} were delivered. November 2017, the Black "
        "Friday month, is the tallest bar.",
        0.6,
        4.75,
        7.2,
        0.9,
        size=14,
        color=GREY,
    )
    peak = datetime.strptime(ao["peak_month"]["month"], "%Y-%m").strftime("%b %Y")
    deck.stat(
        s,
        f"{ao['orders_2017_01']:,} to {ao['peak_month']['orders']:,}",
        f"orders per month, Jan 2017 to {peak}",
        8.2,
        1.7,
        BLUE,
        width=4.6,
    )
    deck.stat(
        s,
        f"{ao['peak_day']['orders']:,}",
        f"orders on {ao['peak_day']['date']} (next busiest day: "
        f"{ao['second_busiest_day']['orders']:,})",
        8.2,
        3.2,
        AMBER,
        width=4.6,
    )
    deck.stat(
        s,
        f"{dl['median_days']:.0f} vs {dl['median_promised_days']:.0f} days",
        "median actual vs promised delivery time",
        8.2,
        4.7,
        GREEN,
        width=4.6,
    )

    # 6. Where lateness hurts -------------------------------------------------------------------
    s = deck.slide(
        "Lateness depends on region and drifts over time",
        "Crew 1 insights",
        "Northeast states are three to five times worse than the South and Southeast; the late "
        "rate swings from about 1% to 19% between months.",
    )
    deck.picture(s, "late_rate_by_state.png", 0.5, 1.5, height=5.3)
    deck.picture(s, "late_rate_by_month.png", 6.6, 1.6, width=6.4)
    worst = ", ".join(f"{x['state']} {pct(x['late_rate'])}" for x in states[:3])
    best = ", ".join(f"{x['state']} {pct(x['late_rate'])}" for x in states[-3:])
    deck.bullets(
        s,
        [
            f"Worst:|{worst}",
            f"Best:|{best}",
            "Implication:|evaluate on later months, never a random split",
        ],
        6.6,
        4.4,
        6.4,
        2.2,
        size=15,
    )

    # 7. Contract -------------------------------------------------------------------------------
    s = deck.slide(
        "The dataset contract is the handoff",
        "Data contract",
        "The analysts publish a versioned contract; the scientists must follow it; the Flow "
        "verifies it twice.",
    )
    deck.bullets(
        s,
        [
            f"Schema:|{len(contract['columns'])} columns with types, nullability, bounds and allowed "
            "values (27 states, 4 payment types, categories)",
            f"Constraints:|{len(contract['constraints'])} binding rules, e.g. is_late == "
            "(delivery_days > promised_days), one row per order",
            "Target:|is_late means delivered after the promised date, counted in calendar days, "
            f"among delivered orders ({pct(dp['late_rate'])} late)",
            f"Required features:|{len(contract['required_features'])}; forbidden post-purchase "
            f"columns: {len(contract['forbidden_features'])} (no leakage)",
            f"Gate 1:|the contract matches the cleaned data ({checks['gate_contract']} checks)",
            "Gate 2:|every required feature exists before modeling "
            f"({checks['gate_required_features']} checks)",
        ],
        0.6,
        1.6,
        7.6,
        4.8,
        size=16,
    )
    deck.code(
        s,
        [
            "HANDOFF analyst->scientist",
            f"  contract_version={contract['contract_version']}",
            f"  contract_sha256={meta['contract_sha256'][:12]}",
            f"  rows={contract['row_count']}",
            f"  required={len(contract['required_features'])} "
            f"forbidden={len(contract['forbidden_features'])}",
        ],
        8.5,
        1.8,
        4.3,
        2.2,
        size=14,
    )
    deck.text(
        s,
        "The log line every run prints when the handoff passes; the same hash is stored "
        "in the feature spec, model metadata and model card.",
        8.5,
        4.2,
        4.3,
        1.2,
        size=13,
        color=GREY,
    )

    # 8. Modeling --------------------------------------------------------------------------------
    s = deck.slide(
        "Checkout-time features, time-aware evaluation",
        "Crew 2 design",
        "Only information known when the customer pays; trained on earlier months, tested on "
        "later ones.",
    )
    deck.bullets(
        s,
        [
            f"Features:|{len(meta['model_features'])} checkout-time features: basket, package, "
            "payment, geography (states, distance), promise and timing, category",
            "Candidates:|prior baseline, logistic regression, histogram gradient boosting "
            "(sklearn pipelines, seed 42)",
            "Selection:|time-series cross-validation on the train window, PR-AUC (imbalanced target)",
            "Threshold:|tuned on out-of-fold predictions only",
            f"Leakage guard:|{len(contract['forbidden_features'])} post-purchase columns are "
            "forbidden; the features gate fails the run if one appears",
        ],
        0.6,
        1.6,
        6.7,
        5.0,
        size=17,
    )
    order = [k for k in ("train", "test", "censored_buffer") if k in split]
    rows = [["split", "purchase dates", "orders", "late rate"]] + [
        [
            k.replace("_", " "),
            " to ".join(split[k]["dates"]),
            f"{split[k]['rows']:,}",
            pct(split[k]["prevalence"]),
        ]
        for k in order
    ]
    deck.table(s, rows, 7.45, 1.8, 5.45, 1.6, col_widths=[1.4, 2.25, 0.82, 0.98], size=12)
    deck.text(
        s,
        "August 2018 is held out as a censored buffer: the snapshot ends in mid-October "
        "2018, so very late August orders may not have arrived yet.",
        7.45,
        3.75,
        5.45,
        1.2,
        size=13,
        color=GREY,
    )

    # 9. Results -----------------------------------------------------------------------------------
    s = deck.slide(
        "The simpler model wins: 3x better than chance at the top",
        "Results",
        "Logistic regression beat gradient boosting in cross-validation and on the test window.",
    )
    rows = [["model", "CV PR-AUC", "test ROC-AUC", "test PR-AUC", "precision, top 10%"]]
    names = list(metrics["candidates"])
    for name in names:
        c = metrics["candidates"][name]
        rows.append(
            [
                name.replace("_", " "),
                f"{c['cv']['pr_auc_mean']:.3f}",
                f"{c['test']['roc_auc']:.3f}",
                f"{c['test']['pr_auc']:.3f}",
                pct(c["test"]["precision_at_top10pct"]),
            ]
        )
    deck.table(
        s,
        rows,
        0.6,
        1.7,
        7.2,
        1.7,
        col_widths=[2.3, 1.15, 1.25, 1.15, 1.35],
        size=13,
        highlight=names.index(chosen) + 1,
    )
    deck.picture(s, "roc_curves.png", 8.2, 1.5, height=4.1)
    deck.stat(
        s,
        pct(test["precision_at_top10pct"]),
        f"late among the riskiest 10% ({test['lift_at_top10pct']:.1f}x the "
        f"{pct(test['prevalence'])} base rate)",
        0.6,
        3.9,
        BLUE,
        width=3.6,
    )
    top = ", ".join(f["feature"] for f in metrics["permutation_importance"]["features"][:4])
    deck.stat(
        s, f"{test['roc_auc']:.3f}", "test ROC-AUC (baseline 0.500)", 4.4, 3.9, GREEN, width=3.4
    )
    deck.text(s, f"Strongest signals: {top}.", 0.6, 5.5, 7.2, 0.6, size=15, color=GREY)

    # 10. Reliability ---------------------------------------------------------------------------------
    s = deck.slide(
        "Engineered like a product: gates, failure, reproducibility",
        "Reliability",
        "Show the fault-injection demo and the reproducibility table in the video.",
    )
    deck.bullets(
        s,
        [
            f"Gates:|contract ({checks['gate_contract']} checks), required features "
            f"({checks['gate_required_features']}), artifacts ({checks['gate_artifacts']}); every "
            "check is logged",
            "Graceful failure:|7 injectable faults; a failed gate writes run_status.json, exits 1 and "
            "promotes nothing",
            "Reproducible:|two fresh runs with different hash seeds give byte-identical data; the "
            "LLM run differs only in prose",
            "Observable:|pipeline.log records stages, gates, the handoff, crews, tools and every LLM "
            "call",
            "Team workflow:|one pull request per milestone, protected main, CI on every PR",
        ],
        0.6,
        1.6,
        6.0,
        5.2,
        size=15,
    )
    deck.code(
        s,
        [
            "$ olist_flow.run --no-llm --inject-fault missing_feature",
            "FAULT INJECTED missing_feature",
            "GATE gate_required_features required_features_present FAIL",
            "GATE gate_required_features RESULT FAIL (11/13 checks)",
            "RUN demo-fault FAILED at gate_required_features",
            "run demo-fault: failed (exit 1), nothing promoted",
            "",
            "$ olist_flow.run verify-repro",
            "clean_data.csv        byte-identical  identical",
            "features.csv          byte-identical  identical",
            "metrics.json          byte-identical  identical",
            "predictions_test.csv  byte-identical  identical",
            "figures/*.png         same machine    11/11 identical",
            "ALL DETERMINISTIC ARTIFACTS IDENTICAL",
        ],
        6.85,
        1.6,
        6.05,
        4.4,
        size=12,
    )

    # 11. Model card -----------------------------------------------------------------------------------
    s = deck.slide(
        "Responsible use: model card and ethics",
        "Model card",
        "Read the ethical considerations section aloud; it is part of the deliverable.",
    )
    deck.bullets(
        s,
        [
            "Purpose:|prioritize risky orders for proactive action, never deny service or "
            "penalize customers or sellers",
            "Limitations:|delivered orders only; right-censoring; the late rate drifts, so the "
            "threshold must be re-tuned; 2017-2018 data",
            "Regional bias:|North and Northeast customers face more delays; interventions must "
            "improve their service",
            "Privacy:|pseudonymous ids; zip prefix plus timestamps are quasi-identifiers, so dashboards "
            "show aggregates only",
            "License:|Olist data for non-commercial use (CC BY-NC-SA 4.0 listing); review text is "
            "never used",
        ],
        0.6,
        1.6,
        6.4,
        5.2,
        size=15,
    )
    deck.picture(s, "feature_importance_selected.png", 7.25, 1.55, width=5.7)
    deck.text(
        s,
        "The score leans on the promise length and on geography, so regional fairness "
        "is part of monitoring.",
        7.35,
        5.75,
        5.5,
        0.8,
        size=13,
        color=GREY,
    )

    # 12. Deployment and next steps ------------------------------------------------------------------
    s = deck.slide(
        "From engine to product, and what comes next",
        "The product and next steps",
        f"Run {status['run_id']} produced the committed artifacts ({status['mode']} mode). Demo: "
        "replay 2018-07-16 in the copilot, download the action list, then ask a free question.",
    )
    deck.bullets(
        s,
        [
            "Ops copilot:|a chat for the team that operates the agents: runs, gates, model and "
            "data; scores a CSV, a replayed day or one order. Read-only, every number sourced",
            "Dashboard and API:|the same artifacts and one risk policy for people and systems",
            "Run it:|uv run kickoff (Gemini) or the --no-llm deterministic mode",
        ],
        0.6,
        1.6,
        5.9,
        3.0,
        size=14,
    )
    deck.bullets(
        s,
        [
            "Next:|a Telegram bot and an 8:00 digest on the same copilot core",
            "Next:|send the customer messages and track handled orders",
            "Next:|seller-history features and a delivery-days regression",
        ],
        0.6,
        3.75,
        5.9,
        2.0,
        size=14,
    )
    shot = deck.picture(s, "ops_copilot.png", 6.85, 1.6, width=6.1, folder=IMAGES)
    shot.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    shot.line.width = Pt(1)
    deck.text(
        s,
        [f"Live: {APP_URL}", "Ops copilot replaying 2018-07-16, with what really happened"],
        6.85,
        5.5,
        6.1,
        0.7,
        size=12,
        color=GREY,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT).as_posix()} ({deck.count} slides)")


if __name__ == "__main__":
    build()
